#!/usr/bin/env python3
"""Audit basic PDF-to-PPTX fidelity and editability signals."""

from __future__ import annotations

import argparse
import json
import re
import sys
from difflib import SequenceMatcher
from pathlib import Path


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def pdf_pages(path: Path) -> list[str]:
    try:
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            return [page.get_text("text") for page in doc]
    except ImportError:
        try:
            from pypdf import PdfReader

            return [(page.extract_text() or "") for page in PdfReader(str(path)).pages]
        except ImportError as exc:
            raise RuntimeError("Install PyMuPDF or pypdf to inspect the source PDF") from exc


def pptx_slides(path: Path) -> list[dict]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError("Install python-pptx to inspect the output presentation") from exc

    prs = Presentation(str(path))
    result = []
    slide_area = int(prs.slide_width) * int(prs.slide_height)
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        text_shapes = 0
        pictures = 0
        largest_picture_ratio = 0.0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = getattr(shape, "text", "")
                if normalize(value):
                    text_shapes += 1
                    texts.append(value)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                ratio = (int(shape.width) * int(shape.height)) / slide_area if slide_area else 0
                largest_picture_ratio = max(largest_picture_ratio, ratio)
        result.append(
            {
                "slide": index,
                "text": "\n".join(texts),
                "text_shapes": text_shapes,
                "pictures": pictures,
                "largest_picture_area_ratio": round(largest_picture_ratio, 4),
                "likely_flattened": largest_picture_ratio >= 0.9 and text_shapes == 0,
            }
        )
    return result


def build_report(pdf_text: list[str], slides: list[dict]) -> dict:
    comparisons = []
    for i in range(max(len(pdf_text), len(slides))):
        source = normalize(pdf_text[i]) if i < len(pdf_text) else ""
        output = normalize(slides[i]["text"]) if i < len(slides) else ""
        ratio = SequenceMatcher(None, source, output, autojunk=False).ratio() if source or output else 1.0
        entry = {
            "page": i + 1,
            "source_characters": len(source),
            "output_characters": len(output),
            "text_similarity": round(ratio, 6),
        }
        if i < len(slides):
            entry.update({k: v for k, v in slides[i].items() if k != "text"})
        comparisons.append(entry)

    flattened = [x["page"] for x in comparisons if x.get("likely_flattened")]
    mismatched = [x["page"] for x in comparisons if x["source_characters"] and x["text_similarity"] < 0.995]
    return {
        "pdf_pages": len(pdf_text),
        "pptx_slides": len(slides),
        "counts_match": len(pdf_text) == len(slides),
        "likely_flattened_pages": flattened,
        "text_review_pages": mismatched,
        "pages": comparisons,
        "note": "Rendered visual comparison is still required.",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", type=Path, dest="json_path")
    args = parser.parse_args()

    for path in (args.pdf, args.pptx):
        if not path.is_file():
            parser.error(f"file not found: {path}")

    try:
        report = build_report(pdf_pages(args.pdf), pptx_slides(args.pptx))
    except RuntimeError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    payload = json.dumps(report, ensure_ascii=False, indent=2)
    print(payload)
    if args.json_path:
        args.json_path.write_text(payload + "\n", encoding="utf-8")
    return 0 if report["counts_match"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
