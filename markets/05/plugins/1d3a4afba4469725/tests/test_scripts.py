#!/usr/bin/env python3
"""Regression tests for the deterministic helpers and packaged artifacts."""

import csv
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PLUGIN_ROOT / "skills/spend-management-analysis"
FIXTURES = PLUGIN_ROOT / "tests/fixtures"
EXPECTED = PLUGIN_ROOT / "tests/expected"


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


build_workbook = load_module("build_workbook", SKILL_ROOT / "scripts/build_workbook.py")
parse_ledes = load_module("parse_ledes", SKILL_ROOT / "scripts/parse_ledes.py")


class WorkbookSafetyTests(unittest.TestCase):
    def test_formula_like_text_is_exported_as_literal_text(self):
        payload = {
            "title": "Injection regression",
            "summary": {
                "headline": "=1+1",
                "audience": "internal review",
                "key_numbers": [{"label": "@test", "value": "+1"}],
            },
            "action_plan": [
                {"priority": 1, "action": "-2+3", "why_it_matters": "=cmd", "timeframe": "now"}
            ],
            "findings": [
                {
                    "category": "Staffing Mix",
                    "item": "=1+1",
                    "severity": "high",
                    "driver": "@external",
                    "recommendation": "+SUM(1,1)",
                }
            ],
            "ai_leverage": [],
            "normalized_data": [{
                "entity": "=1+1",
                "=WEBSERVICE(\"https://example.invalid\")": "value",
                "billed_amount": -25.5,
            }],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            workbook_path = temp / "safe.xlsx"
            findings_path = temp / "safe_findings.csv"

            workbook = Workbook()
            build_workbook.build_summary_sheet(workbook, payload)
            build_workbook.build_action_plan_sheet(workbook, payload)
            build_workbook.build_findings_sheet(workbook, payload)
            build_workbook.build_ai_leverage_sheet(workbook, payload)
            build_workbook.build_normalized_data_sheet(workbook, payload)
            workbook.save(workbook_path)
            build_workbook.write_findings_csv(payload, findings_path)

            loaded = load_workbook(workbook_path, data_only=False)
            formula_cells = [
                f"{sheet.title}!{cell.coordinate}"
                for sheet in loaded.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            ]
            self.assertEqual([], formula_cells)
            self.assertEqual("'=1+1", loaded["Summary"]["B4"].value)
            normalized = loaded["Normalized Data"]
            headers = [cell.value for cell in normalized[1]]
            self.assertIn("'=WEBSERVICE(\"https://example.invalid\")", headers)
            billed_column = headers.index("billed_amount") + 1
            self.assertEqual(-25.5, normalized.cell(row=2, column=billed_column).value)

            with findings_path.open(newline="", encoding="utf-8") as handle:
                finding = next(csv.DictReader(handle))
            self.assertEqual("'=1+1", finding["item"])
            self.assertEqual("'+SUM(1,1)", finding["recommendation"])

    def test_empty_ai_not_assessed_state_explains_missing_evidence(self):
        workbook = Workbook()
        payload = {
            "ai_assessment_status": "not_assessed",
            "ai_assessment_note": "Task-level evidence was unavailable.",
            "ai_leverage": [],
        }
        build_workbook.build_ai_leverage_sheet(workbook, payload)
        message = workbook["AI Leverage"]["A2"].value
        self.assertIn("not assessed", message.lower())
        self.assertIn("task-level evidence was unavailable", message.lower())

        legacy_workbook = Workbook()
        build_workbook.build_ai_leverage_sheet(legacy_workbook, {"ai_leverage": []})
        self.assertIn("not assessed", legacy_workbook["AI Leverage"]["A2"].value.lower())

    def test_empty_ai_assessed_state_reports_no_supported_opportunities(self):
        workbook = Workbook()
        payload = {"ai_assessment_status": "assessed", "ai_leverage": []}
        build_workbook.build_ai_leverage_sheet(workbook, payload)
        message = workbook["AI Leverage"]["A2"].value
        self.assertIn("was assessed", message.lower())
        self.assertIn("no supported", message.lower())
        self.assertNotIn("not assessed", message.lower())


class LedesParserTests(unittest.TestCase):
    def test_valid_fixture_reconciles(self):
        rows, skipped = parse_ledes.parse(FIXTURES / "sample_ledes_1998b_export.txt")
        self.assertEqual(23, len(rows))
        self.assertEqual(0, skipped)
        self.assertEqual(Decimal("124774.50"), sum(Decimal(row["billed_amount"]) for row in rows))
        self.assertEqual(1, sum(row["line_type"] == "expense" for row in rows))

    def test_missing_header_row_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "missing-header.txt"
            path.write_text("LEDES1998B[]\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "missing its pipe-delimited header"):
                parse_ledes.parse(path)

    def test_missing_required_headers_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "missing-required.txt"
            path.write_text(
                "LEDES1998B[]\nEXP/FEE/INV_ADJ_TYPE|LINE_ITEM_TOTAL[]\nF|100[]\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "missing required field"):
                parse_ledes.parse(path)

    def test_nonstandard_header_order_is_rejected(self):
        fixture_lines = (FIXTURES / "sample_ledes_1998b_export.txt").read_text(encoding="utf-8").splitlines()
        headers = fixture_lines[1][:-2].split("|")
        headers[0], headers[1] = headers[1], headers[0]

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "reordered-header.txt"
            path.write_text(
                "LEDES1998B[]\n" + "|".join(headers) + "[]\n" + fixture_lines[2] + "\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "required 24-column LEDES 1998B order"):
                parse_ledes.parse(path)

    def test_missing_record_terminator_is_rejected(self):
        fixture_lines = (FIXTURES / "sample_ledes_1998b_export.txt").read_text(encoding="utf-8").splitlines()
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "missing-terminator.txt"
            path.write_text(
                "LEDES1998B[]\n" + fixture_lines[1] + "\n" + fixture_lines[2][:-2] + "\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "no valid data rows"):
                parse_ledes.parse(path)

    def test_invalid_numeric_row_is_rejected(self):
        fixture_lines = (FIXTURES / "sample_ledes_1998b_export.txt").read_text(encoding="utf-8").splitlines()
        header = fixture_lines[1][:-2].split("|")
        fields = fixture_lines[2][:-2].split("|")
        fields[header.index("LINE_ITEM_TOTAL")] = "not-a-number"

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "invalid-number.txt"
            path.write_text(
                "LEDES1998B[]\n" + fixture_lines[1] + "\n" + "|".join(fields) + "[]\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "no valid data rows"):
                parse_ledes.parse(path)

    def test_cli_neutralizes_text_fields_and_preserves_negative_numbers(self):
        fixture_lines = (FIXTURES / "sample_ledes_1998b_export.txt").read_text(encoding="utf-8").splitlines()
        header = fixture_lines[1][:-2].split("|")
        fields = fixture_lines[2][:-2].split("|")
        fields[header.index("LAW_FIRM_ID")] = "@ATTACKER"
        fields[header.index("LINE_ITEM_DESCRIPTION")] = "=WEBSERVICE(\"https://example.invalid\")"
        fields[header.index("LINE_ITEM_NUMBER_OF_UNITS")] = "-1.5"
        fields[header.index("LINE_ITEM_UNIT_COST")] = "200"
        fields[header.index("LINE_ITEM_TOTAL")] = "-300"
        fields[header.index("INVOICE_TOTAL")] = "-300"

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            input_path = temp / "adversarial-ledes.txt"
            output_path = temp / "normalized.csv"
            input_path.write_text(
                "LEDES1998B[]\n" + fixture_lines[1] + "\n" + "|".join(fields) + "[]\n",
                encoding="utf-8",
            )
            result = subprocess.run(
                [sys.executable, str(SKILL_ROOT / "scripts/parse_ledes.py"), str(input_path), str(output_path)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            with output_path.open(newline="", encoding="utf-8") as handle:
                row = next(csv.DictReader(handle))

        self.assertEqual("'@ATTACKER", row["entity"])
        self.assertEqual("'=WEBSERVICE(\"https://example.invalid\")", row["line_item_description"])
        self.assertEqual(Decimal("-1.5"), Decimal(row["hours"]))
        self.assertEqual(Decimal("200"), Decimal(row["rate"]))
        self.assertEqual(Decimal("-300"), Decimal(row["billed_amount"]))


class PackageArtifactTests(unittest.TestCase):
    def test_public_short_description_meets_final_directory_limit(self):
        manifest = json.loads((PLUGIN_ROOT / ".codex-plugin/plugin.json").read_text(encoding="utf-8"))
        description = manifest["interface"]["shortDescription"]
        self.assertIsInstance(description, str)
        self.assertTrue(description.strip())
        self.assertNotRegex(description, r"[\r\n]")
        self.assertLessEqual(len(description), 30)

    def test_submission_case_counts_and_fixtures(self):
        cases = json.loads((PLUGIN_ROOT / "tests/submission-test-cases.json").read_text())
        self.assertGreaterEqual(len(cases["positive_test_cases"]), 5)
        self.assertGreaterEqual(len(cases["negative_test_cases"]), 3)
        for case in cases["positive_test_cases"]:
            for fixture in case.get("fixtures", []):
                self.assertTrue((PLUGIN_ROOT / "tests" / fixture).is_file(), fixture)

    def test_all_expected_workbooks_use_five_tab_contract(self):
        expected_sheets = ["Summary", "Action Plan", "Findings", "AI Leverage", "Normalized Data"]
        for workbook_path in EXPECTED.rglob("*.xlsx"):
            with self.subTest(workbook=workbook_path.name):
                workbook = load_workbook(workbook_path, read_only=True)
                self.assertEqual(expected_sheets, workbook.sheetnames)
                self.assertEqual("Data Status", workbook["Summary"]["A5"].value)
                self.assertIn("synthetic", workbook["Summary"]["B5"].value.lower())

    def test_expected_reports_identify_synthetic_data(self):
        for report_path in EXPECTED.rglob("report.md"):
            with self.subTest(report=report_path.parent.name):
                report = report_path.read_text(encoding="utf-8").lower()
                self.assertIn("synthetic", report)
                self.assertIn("ai leverage", report)


if __name__ == "__main__":
    unittest.main()
