#!/usr/bin/env python3
"""
Build a polished Excel workbook (and a plain CSV of the findings) from a
spend-management-analysis JSON payload.

Why this exists: every run of the skill needs the same five workbook views --
summary, action plan, findings, AI leverage, and normalized data -- styled the
same way each time. Writing that formatting logic inline in the skill instructions
would mean re-deriving it (and getting the formatting subtly wrong) on every
single invocation. This script makes the output deterministic and reusable.

Usage:
    python build_workbook.py <input.json> <output_basename>

Produces:
    <output_basename>.xlsx  -- five-tab review workbook
    <output_basename>_findings.csv -- flat CSV of just the findings table

Input JSON schema:
{
  "title": "Q3 E-Billing Review - QBR Prep",
  "summary": {
    "headline": "One or two sentence headline finding.",
    "audience": "outside counsel conversation | executive/board | internal review",
    "key_numbers": [{"label": "Total billed", "value": "$286,450"}, ...]
  },
  "action_plan": [
    {
      "priority": 1,
      "action": "Require associate-level staffing on discovery-heavy litigation phases in the next engagement letter with Whitfield Bosch.",
      "why_it_matters": "This single staffing pattern drove the largest overrun in the dataset (+153% on M-1005).",
      "timeframe": "Before next engagement letter"
    }
  ],
  "findings": [
    {
      "category": "Staffing Mix" | "Budget Variance" | "Allocation",
      "item": "Whitfield Bosch - M-1005 Document Review",
      "severity": "high" | "medium" | "low",
      "driver": "Plain-language explanation of why this is happening.",
      "recommendation": "Specific action to take."
    }
  ],
  "ai_assessment_status": "assessed",
  "ai_assessment_note": "Task codes and line descriptions were available.",
  "ai_leverage": [
    {
      "work_type": "Document review (M-1005, Whitfield Bosch)",
      "opportunity_level": "high" | "medium" | "low",
      "rationale": "65 hours of first-pass responsiveness/privilege review, a pattern-based task well suited to AI-assisted triage.",
      "caveat": "AI-assisted review still requires attorney QC before reliance, and must go through the firm's confidentiality/AI-use policy."
    }
  ],
  "normalized_data": [
    {"entity": "Whitfield Bosch", "matter_id": "M-1005", ...},
    ...
  ]
}
"""

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1F2A44"
GOLD = "C9A227"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"

SEVERITY_FILL = {
    "high": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
    "medium": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "low": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
HEADER_FONT = Font(color=WHITE, bold=True, size=11)
TITLE_FONT = Font(color=NAVY, bold=True, size=16)
LABEL_FONT = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))


def safe_spreadsheet_value(value):
    """Keep untrusted text literal in Excel and CSV exports.

    Spreadsheet applications may interpret strings beginning with formula
    control characters as formulas. Genuine numeric values remain numeric;
    formula-like strings are prefixed with an apostrophe so they display as
    text when a user opens the generated files.
    """
    if isinstance(value, str):
        candidate = value.lstrip(" \t\r\n")
        if candidate.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def autosize(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_summary_sheet(wb, data):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = safe_spreadsheet_value(data.get("title", "Spend Management Analysis"))
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    summary = data.get("summary", {})
    ws["A3"] = "Audience"
    ws["A3"].font = LABEL_FONT
    ws["B3"] = safe_spreadsheet_value(summary.get("audience", "not specified"))

    ws["A4"] = "Headline"
    ws["A4"].font = LABEL_FONT
    ws["B4"] = safe_spreadsheet_value(summary.get("headline", ""))
    ws["B4"].alignment = WRAP
    ws.merge_cells("B4:D4")
    ws.row_dimensions[4].height = 45

    row = 6
    ws.cell(row=row, column=1, value="Key Numbers").font = LABEL_FONT
    row += 1
    for kn in summary.get("key_numbers", []):
        ws.cell(row=row, column=1, value=safe_spreadsheet_value(kn.get("label", "")))
        ws.cell(row=row, column=2, value=safe_spreadsheet_value(kn.get("value", "")))
        ws.cell(row=row, column=2).font = Font(bold=True)
        row += 1

    autosize(ws, [22, 22, 22, 40])
    return ws


PRIORITY_FILL = {
    1: PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
    2: PatternFill(start_color="FCE8B2", end_color="FCE8B2", fill_type="solid"),
    3: PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
}


def build_action_plan_sheet(wb, data):
    ws = wb.create_sheet("Action Plan")
    headers = ["Priority", "Action", "Why It Matters", "Timeframe"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    actions = sorted(data.get("action_plan", []), key=lambda a: a.get("priority", 99))
    for r, a in enumerate(actions, start=2):
        priority = a.get("priority", r - 1)
        values = [
            priority,
            a.get("action", ""),
            a.get("why_it_matters", ""),
            a.get("timeframe", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=safe_spreadsheet_value(v))
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if c == 1:
                cell.font = Font(bold=True, size=12)
                cell.fill = PRIORITY_FILL.get(priority, PRIORITY_FILL[3])

    ws.freeze_panes = "A2"
    if actions:
        ws.auto_filter.ref = f"A1:D{len(actions) + 1}"
    autosize(ws, [10, 46, 42, 22])
    return ws


OPPORTUNITY_FILL = {
    "high": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
    "medium": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "low": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}


def build_ai_leverage_sheet(wb, data):
    ws = wb.create_sheet("AI Leverage")
    headers = ["Work Type", "Opportunity Level", "Rationale", "Caveat"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    items = data.get("ai_leverage", [])
    for r, item in enumerate(items, start=2):
        level = (item.get("opportunity_level") or "medium").lower()
        values = [
            item.get("work_type", ""),
            level.capitalize(),
            item.get("rationale", ""),
            item.get("caveat", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=safe_spreadsheet_value(v))
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if c == 2:
                cell.fill = OPPORTUNITY_FILL.get(level, OPPORTUNITY_FILL["medium"])

    if not items:
        status = str(data.get("ai_assessment_status", "not_assessed")).strip().lower()
        note = data.get("ai_assessment_note", "")
        if status == "assessed":
            message = (
                "The available task-level evidence was assessed, but no supported "
                "AI leverage opportunities were identified."
            )
        else:
            reason = note or "Task-level evidence was unavailable."
            message = f"AI leverage was not assessed. {reason}"
        cell = ws.cell(row=2, column=1, value=safe_spreadsheet_value(message))
        cell.alignment = WRAP

    ws.freeze_panes = "A2"
    if items:
        ws.auto_filter.ref = f"A1:D{len(items) + 1}"
    autosize(ws, [30, 16, 46, 42])
    return ws


def build_findings_sheet(wb, data):
    ws = wb.create_sheet("Findings")
    headers = ["Category", "Item", "Severity", "Driver", "Recommendation"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    findings = data.get("findings", [])
    for r, f in enumerate(findings, start=2):
        severity = (f.get("severity") or "medium").lower()
        values = [
            f.get("category", ""),
            f.get("item", ""),
            severity.capitalize(),
            f.get("driver", ""),
            f.get("recommendation", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=safe_spreadsheet_value(v))
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if c == 3:
                cell.fill = SEVERITY_FILL.get(severity, SEVERITY_FILL["medium"])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(len(findings) + 1, 1)}"
    autosize(ws, [16, 34, 12, 46, 46])
    return ws


def build_normalized_data_sheet(wb, data):
    rows = data.get("normalized_data", [])
    ws = wb.create_sheet("Normalized Data")
    if not rows:
        ws["A1"] = "No normalized data provided."
        return ws

    headers = list(rows[0].keys())
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=safe_spreadsheet_value(h))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for r, row in enumerate(rows, start=2):
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c, value=safe_spreadsheet_value(row.get(h, "")))
            cell.alignment = WRAP

            if h in {"rate", "billed_amount", "budget_estimate", "collected_amount"}:
                cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
            elif h in {"hours", "adjusted_hours"}:
                cell.number_format = "0.00"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    width_by_header = {
        "entity": 24,
        "matter_id": 14,
        "matter_name": 34,
        "matter_type": 22,
        "invoice_number": 18,
        "line_date": 14,
        "line_type": 18,
        "task_code": 14,
        "activity_code": 14,
        "timekeeper_role": 18,
        "timekeeper_name": 24,
        "hours": 10,
        "adjusted_hours": 12,
        "rate": 14,
        "billed_amount": 16,
        "budget_estimate": 18,
        "collected_amount": 18,
        "line_item_description": 54,
    }
    autosize(ws, [width_by_header.get(h, 16) for h in headers])
    return ws


def write_findings_csv(data, out_path):
    findings = data.get("findings", [])
    fieldnames = ["category", "item", "severity", "driver", "recommendation"]
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in findings:
            writer.writerow({k: safe_spreadsheet_value(row.get(k, "")) for k in fieldnames})


def main():
    if len(sys.argv) != 3:
        print("Usage: python build_workbook.py <input.json> <output_basename>")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    output_basename = sys.argv[2]

    data = json.loads(input_path.read_text(encoding="utf-8"))

    wb = Workbook()
    build_summary_sheet(wb, data)
    build_action_plan_sheet(wb, data)
    build_findings_sheet(wb, data)
    build_ai_leverage_sheet(wb, data)
    build_normalized_data_sheet(wb, data)

    xlsx_path = f"{output_basename}.xlsx"
    csv_path = f"{output_basename}_findings.csv"
    wb.save(xlsx_path)
    write_findings_csv(data, csv_path)

    print(f"Wrote {xlsx_path}")
    print(f"Wrote {csv_path}")


if __name__ == "__main__":
    main()
