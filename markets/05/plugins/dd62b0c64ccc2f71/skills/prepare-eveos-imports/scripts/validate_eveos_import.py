#!/usr/bin/env python3
"""Read-only structural validator for Eveos conference and speaker workbooks."""

import argparse
from collections import Counter
from datetime import date, datetime, time
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


FORMULA_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--conferences", required=True)
    parser.add_argument("--speakers")
    parser.add_argument("--baseline-conferences")
    parser.add_argument("--baseline-speakers")
    parser.add_argument("--config", required=True)
    parser.add_argument("--report")
    return parser.parse_args()


def add_issue(items, code, message, **context):
    issue = {"code": code, "message": message}
    if context:
        issue["context"] = context
    items.append(issue)


def parse_datetime(value, formats):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, str):
        text = value.strip()
        for fmt in formats:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
    return None


def header_map(ws, header_row):
    result = {}
    for cell in ws[header_row]:
        if cell.value is not None:
            result[str(cell.value).strip()] = cell.column
    return result


def require_columns(headers, requested, errors, entity):
    resolved = {}
    for logical, label in requested.items():
        if label not in headers:
            add_issue(errors, "missing_column", f"Colonne absente : {label}", entity=entity, logical=logical)
        else:
            resolved[logical] = headers[label]
    return resolved


def extract_keys(path, sheet_name, header_row, key_label):
    workbook = load_workbook(path, data_only=False, read_only=False)
    if sheet_name not in workbook.sheetnames:
        return None
    ws = workbook[sheet_name]
    headers = header_map(ws, header_row)
    if key_label not in headers:
        return None
    column = headers[key_label]
    return {
        str(ws.cell(row, column).value).strip()
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, column).value not in (None, "")
    }


def scan_formula_errors(path, workbook, errors):
    values = load_workbook(path, data_only=True, read_only=False)
    for ws in workbook.worksheets:
        value_ws = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cached = value_ws[cell.coordinate].value
                    if isinstance(cached, str) and cached in FORMULA_ERRORS:
                        add_issue(errors, "formula_error", f"Formule en erreur : {cached}", sheet=ws.title, cell=cell.coordinate)


def main():
    args = parse_args()
    config = json.loads(Path(args.config).read_text(encoding="utf-8"))
    errors, warnings = [], []
    header_row = int(config.get("header_row", 1))
    separator = config.get("speaker_separator", ";")
    formats = config.get("date_formats", ["%d/%m/%Y %H:%M"])

    conf_path = Path(args.conferences)
    conf_wb = load_workbook(conf_path, data_only=False, read_only=False)
    conf_sheet = config["conference_sheet"]
    if conf_sheet not in conf_wb.sheetnames:
        add_issue(errors, "missing_sheet", f"Feuille absente : {conf_sheet}", entity="conference")
        report = {"valid": False, "errors": errors, "warnings": warnings, "summary": {}}
        finish(report, args.report)
        return 1

    conf_ws = conf_wb[conf_sheet]
    conf_headers = header_map(conf_ws, header_row)
    conf_cols = require_columns(conf_headers, config["conference_columns"], errors, "conference")
    if "key" not in conf_cols:
        report = {"valid": False, "errors": errors, "warnings": warnings, "summary": {}}
        finish(report, args.report)
        return 1

    conference_rows = [
        row for row in range(header_row + 1, conf_ws.max_row + 1)
        if conf_ws.cell(row, conf_cols["key"]).value not in (None, "")
    ]
    keys = [str(conf_ws.cell(row, conf_cols["key"]).value).strip() for row in conference_rows]
    duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
    for key in duplicates:
        add_issue(errors, "duplicate_conference_key", f"Clé conférence dupliquée : {key}", key=key)

    statuses = Counter()
    allowed_statuses = set(config.get("allowed_statuses", []))
    if "status" in conf_cols:
        for row in conference_rows:
            key = str(conf_ws.cell(row, conf_cols["key"]).value).strip()
            status = conf_ws.cell(row, conf_cols["status"]).value
            statuses[str(status)] += 1
            if allowed_statuses and status not in allowed_statuses:
                add_issue(errors, "invalid_status", f"Statut non autorisé : {status}", row=row, key=key)

    event_start = datetime.fromisoformat(config["event_start"]) if config.get("event_start") else None
    event_end = datetime.fromisoformat(config["event_end"]) if config.get("event_end") else None
    max_duration = config.get("max_duration_minutes")
    expected_number_format = config.get("excel_date_number_format")
    allow_blank_dates = set(config.get("allow_blank_dates_for_keys", []))
    if "start" in conf_cols and "end" in conf_cols:
        for row in conference_rows:
            key = str(conf_ws.cell(row, conf_cols["key"]).value).strip()
            raw_start = conf_ws.cell(row, conf_cols["start"]).value
            raw_end = conf_ws.cell(row, conf_cols["end"]).value
            if raw_start in (None, "") and raw_end in (None, "") and key in allow_blank_dates:
                continue
            start = parse_datetime(raw_start, formats)
            end = parse_datetime(raw_end, formats)
            if start is None or end is None:
                add_issue(errors, "invalid_date", "Date absente ou illisible", row=row, key=key)
                continue
            if start > end:
                add_issue(errors, "invalid_date_order", "La fin précède le début", row=row, key=key)
            if event_start and start < event_start:
                add_issue(errors, "start_out_of_range", "Début hors période", row=row, key=key)
            if event_end and end > event_end:
                add_issue(errors, "end_out_of_range", "Fin hors période", row=row, key=key)
            if max_duration is not None and (end - start).total_seconds() / 60 > float(max_duration):
                add_issue(errors, "duration_too_long", "Durée supérieure au maximum autorisé", row=row, key=key, minutes=(end - start).total_seconds() / 60)
            if expected_number_format:
                for logical in ("start", "end"):
                    cell = conf_ws.cell(row, conf_cols[logical])
                    if isinstance(cell.value, (datetime, date)) and cell.number_format != expected_number_format:
                        add_issue(errors, "invalid_excel_date_format", f"Format Excel incorrect : {cell.number_format}", row=row, key=key, column=logical)

    allow_blank_taxonomy = config.get("allow_blank_taxonomy_for_keys", {})
    required_taxonomies = set(config.get("required_taxonomies", []))
    for label, allowed_values in config.get("taxonomies", {}).items():
        if label not in conf_headers:
            add_issue(errors, "missing_taxonomy_column", f"Colonne de taxonomie absente : {label}")
            continue
        column = conf_headers[label]
        allowed = set(allowed_values)
        blank_allowed = set(allow_blank_taxonomy.get(label, []))
        for row in conference_rows:
            key = str(conf_ws.cell(row, conf_cols["key"]).value).strip()
            raw = conf_ws.cell(row, column).value
            if raw in (None, ""):
                if key not in blank_allowed:
                    target = errors if label in required_taxonomies else warnings
                    add_issue(target, "blank_taxonomy", f"Taxonomie vide : {label}", row=row, key=key)
                continue
            for value in [part.strip() for part in str(raw).split(separator) if part.strip()]:
                if value not in allowed:
                    add_issue(errors, "invalid_taxonomy", f"Valeur de taxonomie non autorisée : {value}", row=row, key=key, taxonomy=label)

    speaker_keys = set()
    if args.speakers:
        sp_path = Path(args.speakers)
        sp_wb = load_workbook(sp_path, data_only=False, read_only=False)
        sp_sheet = config["speaker_sheet"]
        if sp_sheet not in sp_wb.sheetnames:
            add_issue(errors, "missing_sheet", f"Feuille absente : {sp_sheet}", entity="speaker")
        else:
            sp_ws = sp_wb[sp_sheet]
            sp_headers = header_map(sp_ws, header_row)
            sp_cols = require_columns(sp_headers, config["speaker_columns"], errors, "speaker")
            if "key" in sp_cols:
                sp_key_list = [
                    str(sp_ws.cell(row, sp_cols["key"]).value).strip()
                    for row in range(header_row + 1, sp_ws.max_row + 1)
                    if sp_ws.cell(row, sp_cols["key"]).value not in (None, "")
                ]
                speaker_keys = set(sp_key_list)
                for key, count in Counter(sp_key_list).items():
                    if count > 1:
                        add_issue(errors, "duplicate_speaker_key", f"Clé speaker dupliquée : {key}", key=key)
            scan_formula_errors(sp_path, sp_wb, errors)

    if "speakers" in conf_cols:
        for row in conference_rows:
            key = str(conf_ws.cell(row, conf_cols["key"]).value).strip()
            raw = conf_ws.cell(row, conf_cols["speakers"]).value
            if raw in (None, ""):
                continue
            refs = [part.strip() for part in str(raw).split(separator) if part.strip()]
            for ref in refs:
                if args.speakers and ref not in speaker_keys:
                    add_issue(errors, "unknown_speaker_reference", f"Référence speaker inconnue : {ref}", row=row, key=key)

    if args.baseline_conferences:
        baseline_keys = extract_keys(args.baseline_conferences, conf_sheet, header_row, config["conference_columns"]["key"])
        if baseline_keys is None:
            add_issue(errors, "invalid_conference_baseline", "Impossible de lire les clés du fichier conférences de référence")
        else:
            for missing in sorted(baseline_keys - set(keys)):
                add_issue(errors, "missing_existing_conference_key", f"Clé conférence existante disparue : {missing}", key=missing)

    if args.baseline_speakers:
        baseline_keys = extract_keys(args.baseline_speakers, config["speaker_sheet"], header_row, config["speaker_columns"]["key"])
        if baseline_keys is None:
            add_issue(errors, "invalid_speaker_baseline", "Impossible de lire les clés du fichier speakers de référence")
        else:
            for missing in sorted(baseline_keys - speaker_keys):
                add_issue(errors, "missing_existing_speaker_key", f"Clé speaker existante disparue : {missing}", key=missing)

    scan_formula_errors(conf_path, conf_wb, errors)
    report = {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "summary": {
            "conference_count": len(conference_rows),
            "conference_key_count": len(set(keys)),
            "speaker_count": len(speaker_keys),
            "statuses": dict(statuses),
        },
    }
    finish(report, args.report)
    return 0 if not errors else 1


def finish(report, report_path):
    output = json.dumps(report, ensure_ascii=False, indent=2)
    print(output)
    if report_path:
        Path(report_path).write_text(output + "\n", encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
